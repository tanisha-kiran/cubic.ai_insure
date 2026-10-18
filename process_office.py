"""
Layer 1 — Office / text document ingestion.

Normalizes Word, Excel, PowerPoint, HTML, JSON and plain-text files into the
same StructuredContent schema used by process_backend.py (PDF) and
process_audio.py, so retrieval, citations and the reasoner treat every format
identically:

    {"doc_id", "filename", "pages": {"page_1": {"spatial_blocks": [...],
                                                 "raw_text_stream": "..."}}}

"Pages" are the natural unit of each format: a sheet for Excel, a slide for
PowerPoint, ~3000 characters for flowing text (Word, HTML, TXT...). These
files have no rendered page image, so the reasoner answers from the text
layer (its fast, strictly grounded path).

Usage:
    python process_office.py path/to/file.docx doc_id
"""

import csv
import io
import json
import os
import re
import sys

PAGE_CHARS = 3000          # target size of one "page" of flowing text
MAX_SHEET_ROWS = 2000      # cap per sheet so a huge export cannot flood the index
MAX_PAGES = 400

SUPPORTED_SUFFIXES = {".docx", ".xlsx", ".xlsm", ".pptx", ".html", ".htm", ".json",
                      ".txt", ".md", ".csv", ".tsv", ".log", ".xml", ".yaml", ".yml"}


def read_text_file(path: str) -> str:
    raw = open(path, "rb").read()
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="replace")


def _block(lines):
    lines = [" ".join(str(l).split()) for l in lines if str(l).strip()]
    return {"bbox": None, "text_lines": lines} if lines else None


def _paginate(blocks, page_chars=PAGE_CHARS):
    """Group consecutive blocks into pages of about page_chars characters."""
    pages, current, size = [], [], 0
    for block in blocks:
        if block is None:
            continue
        length = sum(len(l) for l in block["text_lines"])
        if current and size + length > page_chars:
            pages.append(current)
            current, size = [], 0
        current.append(block)
        size += length
    if current:
        pages.append(current)
    return pages[:MAX_PAGES]


def _split_long_text(text: str, limit: int = 700):
    """Split a long run of text on paragraph/sentence boundaries into blocks <= limit chars."""
    blocks = []
    for paragraph in re.split(r"\n\s*\n|\r\n\s*\r\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) <= limit:
            blocks.append(_block([paragraph]))
            continue
        sentences = re.split(r"(?<=[.!?])\s+", paragraph)
        buffer = ""
        for sentence in sentences:
            while len(sentence) > limit:  # no punctuation at all (logs, minified text)
                blocks.append(_block([sentence[:limit]]))
                sentence = sentence[limit:]
            if buffer and len(buffer) + len(sentence) > limit:
                blocks.append(_block([buffer]))
                buffer = ""
            buffer = f"{buffer} {sentence}".strip()
        if buffer:
            blocks.append(_block([buffer]))
    return blocks


# ---------------------------------------------------------------- readers
def read_docx(path):
    import docx

    document = docx.Document(path)
    blocks = []
    # Walk the body in order so tables stay where they appear.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            blocks.append(_block([f"# {text}" if style.startswith("heading") or style == "title" else text]))
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[" ".join(c.text.split()) for c in row.cells] for row in table.rows]
            blocks.extend(_table_blocks(rows))
    for section in document.sections:  # headers/footers often hold dates, ids, confidentiality notes
        for part in (section.header, section.footer):
            text = " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            if text:
                blocks.append(_block([f"(header/footer) {text}"]))
    return _paginate(blocks)


def _table_blocks(rows, header_rows=1, rows_per_block=8):
    """Rows as 'Column: value' lines so a value stays attached to its column name."""
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    header = rows[0] if len(rows) > header_rows else []
    body = rows[header_rows:] if header else rows
    blocks = []
    for i in range(0, len(body), rows_per_block):
        lines = []
        for row in body[i:i + rows_per_block]:
            if header and len(header) == len(row):
                cells = [f"{h}: {v}" for h, v in zip(header, row) if v.strip()]
            else:
                cells = [v for v in row if v.strip()]
            lines.append("; ".join(cells))
        blocks.append(_block(lines))
    return blocks


def read_xlsx(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pages = []
    for sheet in workbook.worksheets:
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                break
            rows.append(["" if v is None else str(v) for v in row])
        blocks = [_block([f"Sheet: {sheet.title}"])] + _table_blocks(rows)
        pages.append([b for b in blocks if b])
    return pages[:MAX_PAGES]


def read_pptx(path):
    from pptx import Presentation

    pages = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        blocks = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                blocks.append(_block([p.text for p in shape.text_frame.paragraphs]))
            if getattr(shape, "has_table", False) and shape.has_table:
                blocks.extend(_table_blocks([[c.text for c in r.cells] for r in shape.table.rows]))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            blocks.append(_block([f"(speaker notes) {slide.notes_slide.notes_text_frame.text}"]))
        blocks = [b for b in blocks if b]
        if blocks:
            pages.append(blocks)
    return pages[:MAX_PAGES]


def read_html(path):
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(read_text_file(path), "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "form"]):
        tag.decompose()
    blocks = []
    title = soup.title.get_text(strip=True) if soup.title else ""
    if title:
        blocks.append(_block([f"# {title}"]))
    for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote", "table"]):
        if tag.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in tag.find_all("tr")]
            blocks.extend(_table_blocks(rows))
        elif tag.find_parent("table") is None:
            text = tag.get_text(" ", strip=True)
            if text:
                blocks.append(_block([f"# {text}" if tag.name.startswith("h") else text]))
    return _paginate(blocks)


def _flatten_json(value, prefix=""):
    if isinstance(value, dict):
        for key, item in value.items():
            yield from _flatten_json(item, f"{prefix}.{key}" if prefix else str(key))
    elif isinstance(value, list):
        for i, item in enumerate(value[:500]):
            yield from _flatten_json(item, f"{prefix}[{i}]")
    else:
        yield f"{prefix}: {value}"


def read_json(path):
    try:
        data = json.loads(read_text_file(path))
    except json.JSONDecodeError:
        return read_plain(path)
    lines = list(_flatten_json(data))
    blocks = [_block(lines[i:i + 12]) for i in range(0, len(lines), 12)]
    return _paginate(blocks)


def read_delimited(path, delimiter):
    text = read_text_file(path)
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error:
        return read_plain(path)
    pages = _paginate(_table_blocks(rows[:MAX_SHEET_ROWS]))
    return pages


def read_plain(path):
    return _paginate(_split_long_text(read_text_file(path)))


READERS = {
    ".docx": read_docx, ".xlsx": read_xlsx, ".xlsm": read_xlsx, ".pptx": read_pptx,
    ".html": read_html, ".htm": read_html, ".json": read_json,
    ".csv": lambda p: read_delimited(p, ","), ".tsv": lambda p: read_delimited(p, "\t"),
}


def normalize_office_to_structured_content(path, doc_id, filename=None):
    suffix = os.path.splitext(path)[1].lower()
    if not os.path.exists(path):
        print(f"[-] {path} not found")
        return None
    reader = READERS.get(suffix, read_plain)
    try:
        pages = reader(path)
    except Exception as error:  # corrupt or password-protected file
        print(f"[-] Could not read {path}: {error}")
        return None

    structured = {"doc_id": doc_id, "filename": filename or os.path.basename(path),
                  "total_pages": len(pages), "pages": {}}
    for number, blocks in enumerate(pages, 1):
        structured["pages"][f"page_{number}"] = {
            "image_path": None,
            "spatial_blocks": blocks,
            "raw_text_stream": "\n".join(l for b in blocks for l in b["text_lines"]),
        }
    if not structured["pages"]:
        print(f"[-] {path}: no readable text")
        return None

    os.makedirs("ocr_cache", exist_ok=True)
    with open(f"ocr_cache/{doc_id}_structured.json", "w", encoding="utf-8") as f:
        json.dump(structured, f, ensure_ascii=False, indent=2)
    print(f"[ok] {suffix} read: {len(structured['pages'])} pages, "
          f"{sum(len(p['raw_text_stream']) for p in structured['pages'].values())} chars")
    return structured


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python process_office.py <file> <doc_id>")
        sys.exit(1)
    normalize_office_to_structured_content(sys.argv[1], sys.argv[2])
